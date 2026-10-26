#!/usr/bin/env python3
"""Ingest the production backlog workbook into a flat JSON queue.

Reads backlog/book_videos_backlog.xlsx (the 'Backlog' sheet) and writes
backlog/videos.json: one object per upload. This is the single machine-readable
source of truth for the pipeline; every tool downstream keys off a video `code`.

Run:  python tools/ingest.py   (or: make ingest)

Validates on the way out: 138 entries, codes unique, every CL book has an A and a B.
Exits non-zero if any check fails, so CI can gate on it.
"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "backlog" / "book_videos_backlog.xlsx"
OUT = ROOT / "backlog" / "videos.json"

EXPECTED_COUNT = 138


def parse_range(value: str | None) -> list[int] | None:
    """'8-12 min' -> [8, 12]; '1,200-1,800' -> [1200, 1800]; None -> None."""
    if value is None:
        return None
    nums = re.findall(r"[\d,]+", str(value))
    nums = [int(n.replace(",", "")) for n in nums if n.strip(",")]
    if len(nums) >= 2:
        return [nums[0], nums[1]]
    if len(nums) == 1:
        return [nums[0], nums[0]]
    return None


def parse_text_status(value: str | None) -> str:
    """Workbook phrasing -> the pipeline flag the lint step reads."""
    s = (value or "").lower()
    if s.startswith("public domain"):
        return "public_domain"
    if s.startswith("in copyright"):
        return "copyright"
    raise ValueError(f"Unrecognised text status: {value!r}")


def parse_phase(value: str | None) -> int | None:
    """'Phase 1: public-domain classics (...)' -> 1."""
    m = re.search(r"Phase\s+(\d+)", str(value or ""))
    return int(m.group(1)) if m else None


def load_rows() -> list[dict]:
    if not WORKBOOK.exists():
        sys.exit(f"ERROR: workbook not found: {WORKBOOK}")
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb["Backlog"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col = {name: i + 1 for i, name in enumerate(header)}

    def cell(row: int, name: str):
        return ws.cell(row, col[name]).value

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        code = cell(r, "Video code")
        if not code:
            continue
        rows.append(
            {
                "code": str(code).strip(),
                "book_code": str(cell(r, "Book code")).strip(),
                "upload": (str(cell(r, "Upload")).strip() or None),
                "book": str(cell(r, "Book")).strip(),
                "author": str(cell(r, "Author")).strip(),
                "year": cell(r, "Year"),
                "lane": str(cell(r, "Lane")).strip(),
                "series": (str(cell(r, "Series / group")).strip()
                           if cell(r, "Series / group") else None),
                "text_status": parse_text_status(cell(r, "Text status")),
                "visual_system": (str(cell(r, "Visual system")).strip()
                                  if cell(r, "Visual system") else None),
                "format": str(cell(r, "Format")).strip(),
                "target_minutes": parse_range(cell(r, "Target length")),
                "target_words": parse_range(cell(r, "Target words")),
                "structure": (str(cell(r, "Structure (beat sheet)")).strip()
                              if cell(r, "Structure (beat sheet)") else None),
                "phase": parse_phase(cell(r, "Phase")),
                "week": cell(r, "Week"),
                "status": (str(cell(r, "Status")).strip()
                           if cell(r, "Status") else "Planned"),
            }
        )
    return rows


def validate(rows: list[dict]) -> list[str]:
    errors: list[str] = []
    if len(rows) != EXPECTED_COUNT:
        errors.append(f"expected {EXPECTED_COUNT} entries, got {len(rows)}")

    codes = [r["code"] for r in rows]
    dups = {c for c in codes if codes.count(c) > 1}
    if dups:
        errors.append(f"duplicate codes: {sorted(dups)}")

    cl_uploads: dict[str, set] = defaultdict(set)
    for r in rows:
        if r["book_code"].startswith("CL"):
            cl_uploads[r["book_code"]].add(r["upload"])
    for book_code, uploads in sorted(cl_uploads.items()):
        if uploads != {"A", "B"}:
            errors.append(f"{book_code} should have A and B, has {sorted(uploads)}")

    for r in rows:
        if r["target_minutes"] is None:
            errors.append(f"{r['code']}: missing target_minutes")
        if r["target_words"] is None:
            errors.append(f"{r['code']}: missing target_words")
    return errors


def main() -> int:
    rows = load_rows()
    errors = validate(rows)
    OUT.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(rows)} entries)")
    if errors:
        print("\nVALIDATION FAILED:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        return 1
    print("validation OK: 138 entries, codes unique, every CL book has A and B")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
